import os
import openpyxl

source = "/workspaces/sae-crm-source"

print("\n========================================")
print("SAE CRM — EXCEL SOURCE INVENTORY")
print("========================================")

for root, dirs, files in os.walk(source):
    for filename in sorted(files):
        if filename.lower().endswith((".xlsx", ".xlsm")):
            filepath = os.path.join(root, filename)
            print(f"\nFILE: {filename}")
            print(f"PATH: {filepath}")

            try:
                wb = openpyxl.load_workbook(filepath, read_only=True, data_only=False)

                print(f"SHEETS: {len(wb.sheetnames)}")

                for ws in wb.worksheets:
                    print(f"\n  SHEET: {ws.title}")
                    print(f"  DIMENSION: {ws.max_row} rows x {ws.max_column} columns")

                    shown = 0

                    for row in ws.iter_rows(values_only=True):
                        if any(v is not None and v != "" for v in row):
                            print("  ROW:", repr(row))
                            shown += 1
                            if shown >= 5:
                                break

                wb.close()

            except Exception as e:
                print(f"ERROR: {e}")

print("\n========================================")
print("END OF INVENTORY")
print("========================================")

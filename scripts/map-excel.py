import os
import openpyxl

source = "/workspaces/sae-crm-source"

files = []
for root, dirs, names in os.walk(source):
    for name in names:
        if name.lower().endswith((".xlsx", ".xlsm")):
            files.append(os.path.join(root, name))

for filepath in sorted(files):
    print("\n" + "="*90)
    print("FILE:", os.path.basename(filepath))
    print("="*90)

    wb = openpyxl.load_workbook(filepath, data_only=False, read_only=False)

    for ws in wb.worksheets:
        print("\n--- SHEET:", ws.title, "---")
        print("Dimension:", ws.max_row, "x", ws.max_column)

        print("\nMERGED RANGES:")
        for r in list(ws.merged_cells.ranges)[:30]:
            print(" ", r)

        print("\nNON-EMPTY ROWS:")
        for row_num in range(1, ws.max_row + 1):
            values = [ws.cell(row_num, col).value for col in range(1, ws.max_column + 1)]
            if any(v is not None and v != "" for v in values):
                print(f"ROW {row_num}: {values}")

    wb.close()

print("\nDONE")

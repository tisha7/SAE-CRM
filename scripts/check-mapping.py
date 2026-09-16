import openpyxl, os

source="/workspaces/sae-crm-source/inspect_data/tisha nibe"

files=[
    ("5.SO.xlsx", ["This Year","Marking Condition","LY"]),
    ("6.TSO.xlsx", ["This Year","Condition"]),
    ("Monitoring Templete.xlsx", ["Ach. of SO","Ach of SR","Q1 Ach. SO","Ach of TSO"])
]

for filename,sheets in files:
    path=os.path.join(source,filename)
    wb=openpyxl.load_workbook(path,data_only=False)
    print("\n"+"="*80)
    print(filename)
    print("="*80)

    for sheet in sheets:
        ws=wb[sheet]
        print(f"\n[{sheet}]")

        for r in range(1,ws.max_row+1):
            vals=[ws.cell(r,c).value for c in range(1,min(ws.max_column,22)+1)]
            non=[v for v in vals if v is not None and v!=""]
            if not non:
                continue

            text=" | ".join(str(v) for v in non)

            if (
                r<=8
                or "KPI" in text
                or "Behavioral" in text
                or "Total" in text
                or "Zone-" in text
                or "NPH" in text
                or "SL #" in text
            ):
                print(f"ROW {r}: {text}")

    wb.close()

print("\nMAPPING CHECK COMPLETE")
